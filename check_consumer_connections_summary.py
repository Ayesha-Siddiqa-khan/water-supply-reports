"""Self-check for the Consumer Report category summary exports.

Run: python check_consumer_connections_summary.py
"""

import csv
import io
import json

from openpyxl import load_workbook
from pypdf import PdfReader

from app import app


ROWS = [
    {"serial": 1, "sector": "Domestic", "locality": "", "rate": 4800, "closed": 2, "suspended": 1, "active": 10, "total": 13, "budget": 48000},
    {"serial": 2, "sector": "Private Societies", "locality": "", "rate": 9600, "closed": 1, "suspended": 0, "active": 5, "total": 6, "budget": 48000},
    {"serial": 3, "sector": "Commercial", "locality": "", "rate": "Mixed", "closed": 3, "suspended": 2, "active": 7, "total": 12, "budget": 70000},
]


def main():
    payload = {"summary_rows": ROWS, "grand_total": {}, "sector_totals": {}}
    client = app.test_client()
    query = "?tab=connections&cols=sector,active,total"

    page_response = client.get("/consumer-report")
    assert page_response.status_code == 200
    assert b' id="connections-summary-table"' in page_response.data

    csv_response = client.post("/consumer-report/export/csv" + query, data={"summary_data": json.dumps(payload)})
    assert csv_response.status_code == 200
    csv_rows = list(csv.reader(io.StringIO(csv_response.get_data(as_text=True))))
    assert csv_rows[0] == ["Classification", "Active", "Total Connections"]
    assert [row[0] for row in csv_rows[1:4]] == ["Domestic", "Private Societies", "Commercial"]
    assert csv_rows[-1] == ["GRAND TOTAL", "22", "31"]

    xlsx_response = client.post("/consumer-report/export/xlsx" + query, data={"summary_data": json.dumps(payload)})
    workbook = load_workbook(io.BytesIO(xlsx_response.data), read_only=True)
    assert workbook.sheetnames == ["Details of Connections"]

    pdf_response = client.post("/consumer-report/export/pdf" + query, data={"summary_data": json.dumps(payload)})
    assert pdf_response.status_code == 200
    assert pdf_response.mimetype == "application/pdf"
    assert "details_of_connections.pdf" in pdf_response.headers["Content-Disposition"]
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf_response.data)).pages)
    assert "Details of Connections" in pdf_text
    assert pdf_text.index("Domestic") < pdf_text.index("Private Societies") < pdf_text.index("Commercial")

    print("Consumer connection summary export checks passed.")


if __name__ == "__main__":
    main()
