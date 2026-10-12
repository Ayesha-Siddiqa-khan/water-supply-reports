"""Monthly Account Branch Report.

Provides interactive column and row selection, file upload (CSV / XLSX),
and PDF, Excel, and CSV export for DDO Monthly Account Receipts.
"""

from __future__ import annotations

import csv
import io
import json
import os
import re
import sys
from datetime import datetime

import pandas as pd
from flask import Blueprint, Response, flash, redirect, render_template, request, url_for
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

monthly_account_branch_bp = Blueprint("monthly_account_branch", __name__)


def _app_data_dir() -> str:
    if os.environ.get("VERCEL") == "1":
        return os.path.join("/tmp", "water-supply-report")
    if getattr(sys, "frozen", False):
        return os.path.abspath(os.path.dirname(sys.executable))
    return os.path.abspath(os.path.dirname(__file__))


def _resource_path(*parts: str) -> str:
    if getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS"):
        return os.path.join(sys._MEIPASS, *parts)
    return os.path.join(os.path.abspath(os.path.dirname(__file__)), *parts)


BASE_DIR = _app_data_dir()
UPLOAD_FOLDER = os.path.join(BASE_DIR, "uploads")
CACHE_FILE = os.path.join(UPLOAD_FOLDER, "monthly_account_branch_cache.json")
SAMPLE_CSV = _resource_path("seed", "monthly_account_branch_sample.csv")

NUMERIC_COL_NAMES = {
    "Estimated Budget for the Year",
    "Previous Months Receipts",
    "Actual for Month",
    "Less: OSR Refund",
    "Progressive",
}


def parse_number(val) -> float:
    if val is None or val == "":
        return 0.0
    val_str = str(val).replace(",", "").strip()
    try:
        return float(val_str)
    except Exception:
        return 0.0


def fmt_amt(val) -> str:
    if val is None or val == "":
        return "-"
    try:
        n = float(str(val).replace(",", "").strip())
        if n == 0:
            return "0"
        return f"{int(round(n)):,}"
    except Exception:
        return str(val)


def is_default_water_row(row_dict: dict, row_no: str, row_idx: int) -> bool:
    """Check if row matches the 3 default Water Supply codes/descriptions:
    - C0388047: Receipts on account of sale of water - residential
    - C0388048: Receipts on account of sale of water - commercial
    - C0388093: Arrears of Water Rate (residential)
    """
    code = (str(row_dict.get("Object Code") or "")).strip().upper().replace("O", "0")
    if code.startswith("0C"):
        code = code[1:]
    target_codes = {"C0388047", "C0388048", "C0388093"}
    for tc in target_codes:
        if tc in code or code == tc:
            return True

    head = (str(row_dict.get("Detail Receipt Head") or "")).lower()
    if any(k in head for k in ("sale of water", "arrear of water", "arrears of water")):
        return True

    clean_no = str(row_no or row_dict.get("No.") or "").strip()
    if clean_no in {"11", "12", "18"}:
        return True
    if row_idx in {10, 11, 17}:
        return True
    return False


def parse_csv_content(content_str: str) -> tuple[list[str], list[dict]]:
    reader = csv.reader(io.StringIO(content_str))
    raw_rows = [r for r in reader if any(cell.strip() for cell in r)]
    if not raw_rows:
        return [], []

    headers = [h.strip() for h in raw_rows[0]]
    data_rows: list[dict] = []

    for r in raw_rows[1:]:
        cells = [r[i].strip() if i < len(r) else "" for i in range(len(headers))]
        # Skip trailing summary / grand total rows in the CSV so they are calculated dynamically
        is_total_row = any(
            c.lower() in ("total", "grand total", "sub total") for c in cells[:3]
        )
        if is_total_row:
            continue

        row_dict = {h: cells[i] for i, h in enumerate(headers)}
        idx = len(data_rows)
        data_rows.append(
            {
                "idx": idx,
                "is_default": is_default_water_row(row_dict, row_dict.get("No.", ""), idx),
                "cells": row_dict,
            }
        )
    return headers, data_rows


def parse_excel_bytes(file_bytes: bytes) -> tuple[list[str], list[dict]]:
    df = pd.read_excel(io.BytesIO(file_bytes), header=0)
    headers = [str(c).strip() for c in df.columns]
    data_rows: list[dict] = []
    for _, row in df.iterrows():
        row_dict = {}
        for h in headers:
            val = row.get(h)
            if pd.isna(val):
                row_dict[h] = ""
            elif isinstance(val, float) and val.is_integer():
                row_dict[h] = str(int(val))
            else:
                row_dict[h] = str(val).strip()

        # Check total row
        first_few = [str(row_dict.get(h, "")).lower() for h in headers[:3]]
        if any(w in ("total", "grand total") for w in first_few):
            continue

        idx = len(data_rows)
        data_rows.append(
            {
                "idx": idx,
                "is_default": is_default_water_row(row_dict, row_dict.get("No.", ""), idx),
                "cells": row_dict,
            }
        )
    return headers, data_rows


DEFAULT_SAMPLE_CSV_TEXT = """No.,Object Code,Detail Receipt Head,Estimated Budget for the Year,Previous Months Receipts,Actual for Month,Less: OSR Refund,Progressive
1,B01302,UIPT / Urban Immovable Property tax Share of net proceeds assigned to Districts/TMAs/LGs etc.,25000000,2977000,,,2977000
2,B01313,TTIP / Tax on Transfer of Immovable Property,60000000,5839587,4755253,,10594840
3,C0368301,Grant from Provincial Government PFC (Non Development),230000000,37064000,18532000,,55596000
4,C0368307,Grants from World Bank/Finance Department for PICP,750000000,,,,
5,C0368308,Grant from Provincial Government (Conditional),50000000,,,,
6,C0368309,20% Co-Financing Share for PICP,150000000,,,,
7,C03684,Grant from Provincial Government (Performance/Special),76000000,11388000,,,11388000
8,C03686,Grant from District Government,190000000,,,,
9,C0388016,General Bus Stand fee,9800000,3266078,855275,,4121353
10,C0388042,Fee for slaughtering of animals,901000,299635,78635,,378270
11,C0388047,Receipts on account of sale of water - residential,48100800,1807503,3429940,,5237443
12,C0388048,Receipts on account of sale of water - commercial,3732000,12000,388500,,400500
13,C0388054,Sewerage/Drainage Fee/ Charges/ Nali Tax,9656000,,,,
14,C0388056,Receipts from public latrines,200000,61700,17500,,79200
15,C0388063,"Fee for fairs, agriculture shows, industrial exhibitions, tournaments and other public events",15000,,,,
16,C0388076,Advertisement Fee on billboards/ hoardings,3000000,43400,243000,,286400
17,C0388091,Others Miscellaneous Fee,600000,,47897,,47897
18,C0388093,Arrears of Water Rate (residential),27406841,3062389,,,3062389
19,E01110,Recoveries of expenditure (Cancel Cheque),2000000,,,,
20,E01305,Receipts from investment of Cash Balance by Local Govts (Bank Profit),8000000,,,,
,,Total,1644411641,65821292,28348000,0,94169292"""


def load_cached_or_default_data() -> tuple[list[str], list[dict], str]:
    """Loads parsed data from cache, or falls back to seed sample CSV."""
    filename = "monthly_account_branch_sample.csv"
    try:
        if os.path.exists(CACHE_FILE):
            with open(CACHE_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                headers = data.get("headers", [])
                rows = data.get("rows", [])
                filename = data.get("filename", filename)
                if headers and rows:
                    return headers, rows, filename
    except Exception:
        pass

    # Load from sample CSV file if present
    try:
        if os.path.exists(SAMPLE_CSV):
            with open(SAMPLE_CSV, "r", encoding="utf-8-sig", errors="replace") as f:
                content = f.read()
            headers, rows = parse_csv_content(content)
            if headers and rows:
                return headers, rows, "monthly_account_branch_sample.csv"
    except Exception:
        pass

    # Guaranteed fallback to embedded sample
    headers, rows = parse_csv_content(DEFAULT_SAMPLE_CSV_TEXT)
    return headers, rows, "monthly_account_branch_sample.csv"


def save_cached_data(headers: list[str], rows: list[dict], filename: str) -> None:
    try:
        os.makedirs(os.path.dirname(CACHE_FILE), exist_ok=True)
        with open(CACHE_FILE, "w", encoding="utf-8") as f:
            json.dump({"headers": headers, "rows": rows, "filename": filename}, f)
    except Exception:
        pass


@monthly_account_branch_bp.route("/monthly-account-branch", methods=["GET"])
def monthly_account_branch():
    headers, rows, filename = load_cached_or_default_data()
    today_str = datetime.now().strftime("%d-%m-%Y")
    default_month_year = "September 2026"

    # Default selected row indices
    default_selected_rows = [r["idx"] for r in rows if r["is_default"]]
    if not default_selected_rows and rows:
        default_selected_rows = [r["idx"] for r in rows]

    return render_template(
        "monthly_account_branch.html",
        headers=headers,
        rows=rows,
        filename=filename,
        default_selected_rows=default_selected_rows,
        today_str=today_str,
        default_month_year=default_month_year,
        numeric_cols=list(NUMERIC_COL_NAMES),
    )


@monthly_account_branch_bp.route("/monthly-account-branch/upload", methods=["POST"])
def monthly_account_branch_upload():
    if "file" not in request.files:
        flash("No file selected.")
        return redirect(url_for("monthly_account_branch.monthly_account_branch"))

    file = request.files["file"]
    if not file or not file.filename:
        flash("No file selected.")
        return redirect(url_for("monthly_account_branch.monthly_account_branch"))

    filename = file.filename
    ext = os.path.splitext(filename)[1].lower()

    try:
        if ext in (".xlsx", ".xls"):
            file_bytes = file.read()
            headers, rows = parse_excel_bytes(file_bytes)
        elif ext == ".csv":
            content_str = file.read().decode("utf-8-sig", errors="replace")
            headers, rows = parse_csv_content(content_str)
        else:
            flash("Only .csv or .xlsx files are supported.")
            return redirect(url_for("monthly_account_branch.monthly_account_branch"))

        if not headers or not rows:
            flash("The uploaded file does not contain valid data rows.")
            return redirect(url_for("monthly_account_branch.monthly_account_branch"))

        save_cached_data(headers, rows, filename)
        flash(f"Successfully loaded {len(rows)} rows from {filename}.")
    except Exception as e:
        flash(f"Failed to process file: {e}")

    return redirect(url_for("monthly_account_branch.monthly_account_branch"))


@monthly_account_branch_bp.route("/monthly-account-branch/reset", methods=["POST"])
def monthly_account_branch_reset():
    try:
        if os.path.exists(CACHE_FILE):
            os.remove(CACHE_FILE)
    except Exception:
        pass
    flash("Reset to default sample water supply report data.")
    return redirect(url_for("monthly_account_branch.monthly_account_branch"))


def _extract_selection_args() -> tuple[list[str], list[int], str, str]:
    headers, rows, _ = load_cached_or_default_data()

    # Selected columns
    raw_cols = request.form.get("selected_cols") or request.args.get("selected_cols") or ""
    if raw_cols:
        try:
            selected_cols = json.loads(raw_cols) if raw_cols.startswith("[") else [c.strip() for c in raw_cols.split(",") if c.strip()]
        except Exception:
            selected_cols = [c.strip() for c in raw_cols.split(",") if c.strip()]
    else:
        selected_cols = headers

    # Keep only columns that exist
    selected_cols = [c for c in selected_cols if c in headers]
    if not selected_cols:
        selected_cols = headers

    # Selected rows
    raw_rows = request.form.get("selected_rows") or request.args.get("selected_rows") or ""
    if raw_rows:
        try:
            row_items = json.loads(raw_rows) if raw_rows.startswith("[") else [int(x.strip()) for x in raw_rows.split(",") if x.strip().isdigit()]
            selected_row_indices = [int(x) for x in row_items]
        except Exception:
            selected_row_indices = [int(x.strip()) for x in raw_rows.split(",") if x.strip().isdigit()]
    else:
        # Default water rows
        selected_row_indices = [r["idx"] for r in rows if r["is_default"]]
        if not selected_row_indices:
            selected_row_indices = [r["idx"] for r in rows]

    month_year = (request.form.get("month_year") or request.args.get("month_year") or "September 2026").strip()
    print_date = (request.form.get("print_date") or request.args.get("print_date") or datetime.now().strftime("%d-%m-%Y")).strip()

    return selected_cols, selected_row_indices, month_year, print_date


def generate_monthly_account_branch_pdf(
    headers: list[str],
    rows: list[dict],
    selected_cols: list[str],
    selected_row_indices: list[int],
    month_year: str = "September 2026",
    print_date: str = "",
) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        topMargin=8 * mm,
        bottomMargin=8 * mm,
        leftMargin=8 * mm,
        rightMargin=8 * mm,
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "MABTitle",
        parent=styles["Heading1"],
        fontSize=16,
        textColor=colors.black,
        alignment=1,
        spaceAfter=3 * mm,
        fontName="Helvetica-Bold",
    )
    sub_style = ParagraphStyle(
        "MABSub",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        textColor=colors.black,
        alignment=1,
        spaceAfter=4 * mm,
    )
    th_style = ParagraphStyle(
        "MABTh",
        parent=styles["Normal"],
        fontSize=10,
        leading=13,
        fontName="Helvetica-Bold",
        textColor=colors.black,
        alignment=1,
    )
    td_style = ParagraphStyle(
        "MABTd",
        parent=styles["Normal"],
        fontSize=10,
        leading=13,
        fontName="Helvetica",
        alignment=0,
        textColor=colors.black,
    )
    td_c_style = ParagraphStyle(
        "MABTdC",
        parent=styles["Normal"],
        fontSize=10,
        leading=13,
        fontName="Helvetica",
        alignment=1,
        textColor=colors.black,
    )
    td_r_style = ParagraphStyle(
        "MABTdR",
        parent=styles["Normal"],
        fontSize=10,
        leading=13,
        fontName="Helvetica",
        alignment=2,
        textColor=colors.black,
    )
    grand_label_style = ParagraphStyle(
        "MABGrandLabel",
        parent=styles["Normal"],
        fontSize=10,
        leading=13,
        fontName="Helvetica-Bold",
        alignment=0,
        textColor=colors.black,
    )
    grand_r_style = ParagraphStyle(
        "MABGrandR",
        parent=styles["Normal"],
        fontSize=10,
        leading=13,
        fontName="Helvetica-Bold",
        alignment=2,
        textColor=colors.black,
    )

    elements = [
        Paragraph("Summary Monthly Account Receipts DDO", title_style),
        Paragraph(
            f"<b>For the Month of:</b> {month_year} &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; "
            f"<b>Print Date:</b> {print_date}",
            sub_style,
        ),
        Spacer(1, 2 * mm),
    ]

    # Dynamically allocate column widths across 281mm printable width
    page_w = 281 * mm
    base_weights = {
        "No.": 12,
        "Object Code": 28,
        "Detail Receipt Head": 75,
        "Estimated Budget for the Year": 36,
        "Previous Months Receipts": 34,
        "Actual for Month": 30,
        "Less: OSR Refund": 30,
        "Progressive": 36,
    }
    active_weights = [base_weights.get(c, 30) for c in selected_cols]
    tot_weight = sum(active_weights) or 1
    col_widths = [(w / tot_weight) * page_w for w in active_weights]

    header_cells = [Paragraph(f"<b>{c}</b>", th_style) for c in selected_cols]
    data_table = [header_cells]

    # Filter chosen rows
    chosen_rows = [r for r in rows if r["idx"] in selected_row_indices]
    col_sums = {c: 0.0 for c in selected_cols if c in NUMERIC_COL_NAMES}

    for r in chosen_rows:
        row_cells = []
        row_cells_data = r.get("cells") or r.get("values") or {}
        for c in selected_cols:
            raw_v = row_cells_data.get(c, "")
            if c in NUMERIC_COL_NAMES:
                num = parse_number(raw_v)
                col_sums[c] += num
                disp = fmt_amt(num) if (raw_v != "" or num != 0) else "-"
                row_cells.append(Paragraph(disp, td_r_style))
            elif c == "No.":
                row_cells.append(Paragraph(str(raw_v or "-"), td_c_style))
            elif c == "Object Code":
                row_cells.append(Paragraph(str(raw_v or "-"), td_c_style))
            else:
                row_cells.append(Paragraph(str(raw_v or "-"), td_style))
        data_table.append(row_cells)

    # Grand total row
    grand_cells = []
    total_labeled = False
    for c in selected_cols:
        if c in NUMERIC_COL_NAMES:
            grand_cells.append(Paragraph(f"<b>{fmt_amt(col_sums[c])}</b>", grand_r_style))
        else:
            if not total_labeled and c in ("Detail Receipt Head", "Object Code", "No."):
                grand_cells.append(
                    Paragraph(f"<b>Total ({len(chosen_rows)} Items)</b>", grand_label_style)
                )
                total_labeled = True
            else:
                grand_cells.append(Paragraph("", td_c_style))
    if not total_labeled and grand_cells:
        grand_cells[0] = Paragraph(f"<b>Total ({len(chosen_rows)} Items)</b>", grand_label_style)
    data_table.append(grand_cells)

    t = Table(data_table, colWidths=col_widths, repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f0f0")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("BACKGROUND", (0, 1), (-1, -2), colors.white),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#e0e0e0")),
                ("LINEABOVE", (0, -1), (-1, -1), 1.2, colors.black),
            ]
        )
    )
    elements.append(t)
    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


@monthly_account_branch_bp.route("/monthly-account-branch/export/pdf", methods=["GET", "POST"])
def monthly_account_branch_export_pdf():
    headers, rows, _ = load_cached_or_default_data()
    selected_cols, selected_row_indices, month_year, print_date = _extract_selection_args()

    pdf_bytes = generate_monthly_account_branch_pdf(
        headers=headers,
        rows=rows,
        selected_cols=selected_cols,
        selected_row_indices=selected_row_indices,
        month_year=month_year,
        print_date=print_date,
    )

    clean_my = re.sub(r"[^a-zA-Z0-9]+", "_", month_year).strip("_")
    filename = f"monthly_account_branch_report_{clean_my}.pdf"
    return Response(
        pdf_bytes,
        mimetype="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@monthly_account_branch_bp.route("/monthly-account-branch/export/excel", methods=["GET", "POST"])
def monthly_account_branch_export_excel():
    headers, rows, _ = load_cached_or_default_data()
    selected_cols, selected_row_indices, month_year, print_date = _extract_selection_args()

    chosen_rows = [r for r in rows if r["idx"] in selected_row_indices]
    col_sums = {c: 0.0 for c in selected_cols if c in NUMERIC_COL_NAMES}

    excel_rows = []
    for r in chosen_rows:
        row_dict = {}
        row_data = r.get("cells") or r.get("values") or {}
        for c in selected_cols:
            val = row_data.get(c, "")
            if c in NUMERIC_COL_NAMES:
                num = parse_number(val)
                col_sums[c] += num
                row_dict[c] = num if (val != "" or num != 0) else 0
            else:
                row_dict[c] = val
        excel_rows.append(row_dict)

    # Total row
    total_dict = {}
    total_labeled = False
    for c in selected_cols:
        if c in NUMERIC_COL_NAMES:
            total_dict[c] = col_sums[c]
        else:
            if not total_labeled and c in ("Detail Receipt Head", "Object Code", "No."):
                total_dict[c] = f"Total ({len(chosen_rows)} Items)"
                total_labeled = True
            else:
                total_dict[c] = ""
    if not total_labeled and selected_cols:
        total_dict[selected_cols[0]] = f"Total ({len(chosen_rows)} Items)"
    excel_rows.append(total_dict)

    df = pd.DataFrame(excel_rows, columns=selected_cols)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Monthly Account", index=False, startrow=3)
        ws = writer.sheets["Monthly Account"]

        # Insert Title and Subtitle
        ws.cell(row=1, column=1, value="Summary Monthly Account Receipts DDO")
        ws.cell(
            row=2,
            column=1,
            value=f"For the Month of: {month_year}    |    Print Date: {print_date}",
        )

        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        title_font = Font(name="Calibri", size=14, bold=True)
        sub_font = Font(name="Calibri", size=11, bold=True)
        header_font = Font(name="Calibri", size=11, bold=True)
        total_font = Font(name="Calibri", size=11, bold=True)
        data_font = Font(name="Calibri", size=10)

        gray_fill = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")
        total_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")

        thin_side = Side(border_style="thin", color="000000")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        double_bottom = Border(
            left=thin_side,
            right=thin_side,
            top=thin_side,
            bottom=Side(border_style="double", color="000000"),
        )

        ws.cell(row=1, column=1).font = title_font
        ws.cell(row=2, column=1).font = sub_font

        # Format header row (Row 4)
        for col_idx in range(1, len(selected_cols) + 1):
            c_cell = ws.cell(row=4, column=col_idx)
            c_cell.font = header_font
            c_cell.fill = gray_fill
            c_cell.border = thin_border
            c_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Format data rows (Row 5 to 5 + len(chosen_rows) - 1)
        last_data_row = 4 + len(chosen_rows)
        for r_idx in range(5, last_data_row + 1):
            for col_idx, col_name in enumerate(selected_cols, start=1):
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                if col_name in NUMERIC_COL_NAMES:
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_name in ("No.", "Object Code"):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Format Total row (last row)
        tot_row_idx = last_data_row + 1
        for col_idx, col_name in enumerate(selected_cols, start=1):
            t_cell = ws.cell(row=tot_row_idx, column=col_idx)
            t_cell.font = total_font
            t_cell.fill = total_fill
            t_cell.border = double_bottom
            if col_name in NUMERIC_COL_NAMES:
                t_cell.number_format = "#,##0"
                t_cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                t_cell.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-fit column widths
        for col_idx, col_name in enumerate(selected_cols, start=1):
            col_letter = ws.cell(row=4, column=col_idx).column_letter
            max_len = max(len(str(col_name)), 10)
            for r_idx in range(5, tot_row_idx + 1):
                v = ws.cell(row=r_idx, column=col_idx).value
                if v is not None:
                    max_len = max(max_len, len(str(v)))
            ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

    buf.seek(0)
    clean_my = re.sub(r"[^a-zA-Z0-9]+", "_", month_year).strip("_")
    filename = f"monthly_account_branch_report_{clean_my}.xlsx"
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@monthly_account_branch_bp.route("/monthly-account-branch/export/csv", methods=["GET", "POST"])
def monthly_account_branch_export_csv():
    headers, rows, _ = load_cached_or_default_data()
    selected_cols, selected_row_indices, month_year, _ = _extract_selection_args()

    chosen_rows = [r for r in rows if r["idx"] in selected_row_indices]
    col_sums = {c: 0.0 for c in selected_cols if c in NUMERIC_COL_NAMES}

    out = io.StringIO()
    writer = csv.writer(out)

    writer.writerow(selected_cols)
    for r in chosen_rows:
        row_vals = []
        row_data = r.get("cells") or r.get("values") or {}
        for c in selected_cols:
            val = row_data.get(c, "")
            if c in NUMERIC_COL_NAMES:
                num = parse_number(val)
                col_sums[c] += num
                row_vals.append(str(int(round(num))) if (val != "" or num != 0) else "")
            else:
                row_vals.append(val)
        writer.writerow(row_vals)

    # Total row
    total_vals = []
    total_labeled = False
    for c in selected_cols:
        if c in NUMERIC_COL_NAMES:
            total_vals.append(str(int(round(col_sums[c]))))
        else:
            if not total_labeled and c in ("Detail Receipt Head", "Object Code", "No."):
                total_vals.append(f"Total ({len(chosen_rows)} Items)")
                total_labeled = True
            else:
                total_vals.append("")
    if not total_labeled and selected_cols:
        total_vals[0] = f"Total ({len(chosen_rows)} Items)"
    writer.writerow(total_vals)

    # UTF-8 with BOM for Excel compatibility
    csv_bytes = out.getvalue().encode("utf-8-sig")
    clean_my = re.sub(r"[^a-zA-Z0-9]+", "_", month_year).strip("_")
    filename = f"monthly_account_branch_report_{clean_my}.csv"
    return Response(
        csv_bytes,
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
